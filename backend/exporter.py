from pathlib import Path
from typing import List, Dict, Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADERS = [
    "Balloon No.",
    "Drawing Reference / Text",
    "Type",
    "Nominal Value",
    "Units",
    "Upper Tolerance",
    "Lower Tolerance",
    "Surface Finish",
    "Process",
    "Datum Ref",
    "Tolerance Zone",
    "Quantity",
    "Material",
    "Description",
    "Page No.",
    "Remarks",
]

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
ALT_FILL = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")


def export_to_excel(balloons: List[Dict[str, Any]], output_path: str) -> str:
    """
    Generate an Excel workbook from a list of balloon dicts.

    Args:
        balloons:     list of balloon dicts (from MongoDB, _id already converted to str)
        output_path:  absolute path where the .xlsx file should be saved

    Returns:
        output_path (for convenience)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Balloon Annotations"
    ws.sheet_view.showGridLines = False

    # ── Header row ────────────────────────────────────────────────────────────
    ws.row_dimensions[1].height = 32
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER

    # ── Data rows ─────────────────────────────────────────────────────────────
    sorted_balloons = sorted(balloons, key=lambda b: (b.get("page", 1), b.get("balloon_no", 0)))

    for row_idx, balloon in enumerate(sorted_balloons, start=2):
        unit = balloon.get("units", "")
        if not unit and balloon.get("type") == "Dimension":
            unit = "°" if "°" in str(balloon.get("text", "")) else "mm"

        row_data = [
            balloon.get("balloon_no", ""),
            balloon.get("text", ""),
            balloon.get("type", ""),
            balloon.get("nominal_value", ""),
            unit,
            balloon.get("tolerance_upper", ""),
            balloon.get("tolerance_lower", ""),
            balloon.get("surface_finish", ""),
            balloon.get("process", ""),
            balloon.get("datum_ref", ""),
            balloon.get("tolerance_zone", ""),
            balloon.get("quantity", ""),
            balloon.get("material", ""),
            balloon.get("description", ""),
            balloon.get("page", ""),
            balloon.get("remarks", ""),
        ]
        fill = ALT_FILL if row_idx % 2 == 0 else None
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN_BORDER
            if fill:
                cell.fill = fill
        ws.row_dimensions[row_idx].height = 20

    # ── Auto column widths ────────────────────────────────────────────────────
    for col_idx, header in enumerate(HEADERS, start=1):
        max_len = len(header)
        col_letter = get_column_letter(col_idx)
        for row_idx in range(2, ws.max_row + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        # Cap at 50 to avoid absurdly wide columns
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    # ── Save ──────────────────────────────────────────────────────────────────
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
